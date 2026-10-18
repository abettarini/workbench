"""Fogli costi separati (tavolo base, cassettiera, alzatina) con negozio, prezzo verificato e link,
più il confronto prezzi tra Tecnomat, OBI e Leroy Merlin (ricerca del 18/10/2026)."""
import json
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

F = "Arial"
BLUE, GREEN, NORM, BOLD = Font(name=F, color="0000FF"), Font(name=F, color="008000"), Font(name=F), Font(name=F, bold=True)
H1, NOTE = Font(name=F, bold=True, size=14), Font(name=F, italic=True, color="555555")
LINK = Font(name=F, color="1F5FA8", underline="single")
HEAD_FILL, YEL, SEL = PatternFill("solid", fgColor="E9DFCC"), PatternFill("solid", fgColor="FFF2A8"), PatternFill("solid", fgColor="DCEBD6")
thin = Side(style="thin", color="BBBBBB"); BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
EUR = '#,##0.00 "€";-#,##0.00 "€";"-"'
EURM = '#,##0.00 "€/m"'
DATA = "18/10/2026"
SRC = {"ver": f"verificato {DATA}", "stima": "stima – da verificare", "base": "incluso nel tavolo base"}

d = json.load(open("parts_kvh.json")); cf = json.load(open("confronto.json"))
parts = d["parts"]

def setc(ws, ref, value, font=NORM, fmt=None, fill=None, wrap=False, border=True):
    c = ws[ref]; c.value = value; c.font = font
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    if border: c.border = BOX
    c.alignment = Alignment(wrap_text=wrap, vertical="center")
    return c

def header(ws, row, cols, widths):
    for j, (c, w) in enumerate(zip(cols, widths), 1):
        cell = ws.cell(row=row, column=j, value=c); cell.font = BOLD; cell.fill = HEAD_FILL; cell.border = BOX
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(j)].width = w

def link(ws, ref, url):
    if url:
        c = setc(ws, ref, "apri", LINK); c.hyperlink = url
    else:
        setc(ws, ref, "")

def distinta(ws, title, subtitle, items, cats, foot):
    """items: (categoria, articolo, uso, qtà, unità, prezzo|formula, fonte, collegato, negozio, url, note)"""
    ws["A1"] = title; ws["A1"].font = H1
    ws["A2"] = subtitle; ws["A2"].font = NOTE
    header(ws, 4, ["Categoria", "Articolo (prodotto trovato)", "Uso", "Q.tà", "Unità", "Prezzo unitario", "Totale",
                   "Fonte prezzo", "Negozio", "Link", "Note"], [15, 46, 34, 7, 7, 13, 12, 20, 18, 7, 46])
    r = 5
    for cat, art, use, q, u, pr, src, linked, shop, url, note in items:
        setc(ws, f"A{r}", cat); setc(ws, f"B{r}", art, wrap=True); setc(ws, f"C{r}", use, wrap=True)
        setc(ws, f"D{r}", q, GREEN if linked else BLUE, "0"); setc(ws, f"E{r}", u)
        setc(ws, f"F{r}", pr, GREEN if linked else BLUE, EUR, fill=YEL if src == "stima" else None)
        setc(ws, f"G{r}", f"=D{r}*F{r}", NORM, EUR)
        setc(ws, f"H{r}", SRC[src], wrap=True); setc(ws, f"I{r}", shop, wrap=True); link(ws, f"J{r}", url)
        setc(ws, f"K{r}", note, wrap=True)
        ws.row_dimensions[r].height = 32
        r += 1
    first, last = 5, r - 1; r += 1
    def tot(label, formula, bold=False):
        nonlocal r
        setc(ws, f"F{r}", label, BOLD if bold else NORM, border=False); setc(ws, f"G{r}", formula, BOLD if bold else NORM, EUR); r += 1
    tot("Prezzi verificati", f'=SUMIFS(G{first}:G{last},H{first}:H{last},"{SRC["ver"]}")')
    tot("Stime", f'=SUMIFS(G{first}:G{last},H{first}:H{last},"{SRC["stima"]}")')
    tot("TOTALE", f"=SUM(G{first}:G{last})", True)
    r += 1
    for c in cats: tot(f"di cui {c.lower()}", f'=SUMIFS(G{first}:G{last},A{first}:A{last},"{c}")')
    r += 1
    for shop in sorted({it[8] for it in items}):
        tot(f"da {shop}", f'=SUMIFS(G{first}:G{last},I{first}:I{last},"{shop}")')
    r += 1
    for line in foot:
        ws[f"A{r}"] = line; ws[f"A{r}"].font = NOTE; r += 1
    ws.freeze_panes = "A5"

def dims(p):
    x0, y0, z0, x1, y1, z1 = p["box"]; return sorted([round(x1 - x0), round(y1 - y0), round(z1 - z0)], reverse=True)

TM, OB, LM, FO = "Tecnomat", "OBI", "Leroy Merlin", "ferramenta/online"
U = {  # link prodotti (dalle ricerche del 18/10/2026)
 "kvh": "https://www.tecnomat.it/it/trave-kvh-80x120x400-cm-25102417/",
 "lam40": "https://www.tecnomat.it/it/tavola-abete-lam-200x80-cm-spess-40-25035058/",
 "osb": "https://www.tecnomat.it/it/osb-3-pioppo-spess-18-mm-dim-250x125-cm-10043418/",
 "list2": "https://www.tecnomat.it/it/listello-abete-grezzo-2-m-30x40-mm-25039932/",
 "m10x180": "https://www.obi-italia.it/p/6282354",
 "m10x160": "https://www.obi-italia.it/p/6253421",
 "barra": "https://www.tecnomat.it/it/barra-filettata-zinc-m10x1m-10066674/",
 "spine": "https://www.obi-italia.it/p/4884714",
 "colla": "https://www.tecnomat.it/it/colla-pavimenti-flottanti-tover-500-ml-10036806/",
 "piedi": "https://www.tecnomat.it/it/piede-tondo-nero-o40-2-pz-10038522/",
 "ragno": "https://www.obi-italia.it/p/4611364",
 "squad": "https://www.tecnomat.it/it/squadretta-a-l-30x38x30-per-vite-10027452/",
 "staffa": "https://www.obi-italia.it/p/3786233",
 "dadiaut": "https://www.tecnomat.it/it/dadi-aut-bassi-m10-inox-316-10pz-10055310/",
 "rondlarghe": "https://www.leroymerlin.it/prodotti/8-rondelle-piane-larghe-in-acciaio-zincato-standers-int-10-est-29-48-mm-82583073.html",
 "canalina": "https://www.tecnomat.it/it/canalina-ufficio-60x40-mm-2-scomparti-10052978/",
 "passacavi": "https://www.tecnomat.it/it/bocchette-passacavi-in-pvc-nero-o60-2pz-25019432/",
 "olio": "https://www.tecnomat.it/it/olio-facile-appl-int-legno-adler-750-ml-10070802/",
 "impr": "https://www.tecnomat.it/it/impregnante-acqua-cipir-incolore-0-75-l-10006177/",
 "pio18": "https://www.tecnomat.it/it/pioppo-spess-18-mm-dim-120x60-cm-25016807/",
 "pio15": "https://www.tecnomat.it/it/pioppo-spess-15-mm-120x60-cm-25022690/",
 "fibra3": "https://www.tecnomat.it/it/fibra-grezza-spess-3-mm-dim-250x125-cm-10048894/",
 "guide": "https://www.tecnomat.it/it/guide-p-cas-acc-estraz-tot-h45-500mm-bnc-25060371/",
 "tav18": "https://www.tecnomat.it/it/tavola-abete-lam-100x20-cm-spess-18-10074965/",
}

# ============================================================== 1. TAVOLO BASE
wb = Workbook(); ws = wb.active; ws.title = "Distinta costi"
V, S, B = "ver", "stima", "base"
base_items = [
    ("Legno", "Trave legno massiccio strutturale KVH giuntato 80×120×4000", "gambe, traverse e i 4 montanti dell'alzatina (~2,6 m di scorta)",
     "='Confronto travi'!I8", "pz", "='Confronto travi'!D8", V, True, TM, U["kvh"], "Solo Tecnomat; OBI e Leroy Merlin non hanno KVH"),
    ("Legno", "Tavola abete lamellare 2000×800×40", "piano di lavoro (intero)", 1, "pz", 79.00, V, False, TM, U["lam40"],
     "Leroy Merlin: 147,45 € (marketplace); OBI: non c'è lo spessore 40"),
    ("Legno", "OSB 3 pioppo 18 mm 2500×1250", "3 ripiani bassi + rinforzo ricarica (2 strati)", 1, "pz", 27.50, V, False, TM, U["osb"],
     "OBI 31,49 €, Leroy Merlin 30,50 €"),
    ("Legno", "Listello abete grezzo 30×40×2000", "reggiripiano: 4 × 560 + 2 × 260", 2, "pz", 3.10, V, False, TM, U["list2"],
     "Grezzo va bene: resta nascosto sotto i ripiani"),
    ("Ferramenta telaio", "Vite a testa cilindrica zincata M10×180 (B1)", "traverse laterali", 16, "pz", 1.39, V, False, OB, U["m10x180"],
     "In negozio verifica che sia a esagono incassato (TCEI); TCEI M10×180 non trovate a scaffale altrove"),
    ("Ferramenta telaio", "Viti a testa cilindrica M10×160, conf. 10 pz (B2)", "traverse longitudinali", 1, "conf", 13.09, V, False, OB, U["m10x160"],
     "Accorciale a 140 mm con il seghetto (avvita prima un dado per ripulire il filetto)"),
    ("Ferramenta telaio", "Dado a tamburo M10 (D)", "34 nodi + 4 predisposizione alzatina", 38, "pz", 0.80, S, False, FO, "",
     "Non venduto dai tre negozi: ferramenta, Würth o online. Misura il diametro prima di forare"),
    ("Ferramenta telaio", "Barra filettata M10×1000 zincata (R)", "4 spezzoni da 220 per le gambe centrali", 1, "pz", 1.65, V, False, TM, U["barra"], "OBI 3,19 €, LM 1,75 €"),
    ("Ferramenta telaio", "Spine in faggio zigrinato Ø12, conf. 25 pz", "2 per testata (68)", 3, "conf", 3.69, V, False, OB, U["spine"],
     "Lunghezza non indicata: servono ~50 mm. In alternativa tondo faggio Ø12×1000 Tecnomat 0,80 € da tagliare"),
    ("Ferramenta telaio", "Collante vinilico D3 Tover 500 ml", "spine, rinforzo, listelli", 1, "pz", 4.95, V, False, TM, U["colla"], ""),
    ("Ferramenta telaio", "Piede regolabile Ø40 vite M10, conf. 2 pz", "sotto le gambe", 4, "conf", 1.65, V, False, TM, U["piedi"], ""),
    ("Ferramenta telaio", "Dado ragno M10 zincato (al posto dell'inserto a vite)", "sede filettata sotto ogni gamba", 8, "pz", 2.79, V, False, OB, U["ragno"],
     "Foro Ø12 invece di Ø14; pezzi per confezione non indicati. Inserti M10 a vite: solo marketplace LM"),
    ("Ferramenta telaio", "Squadretta a L 30×38×30 zincata", "fissaggio del piano", 16, "pz", 0.50, V, False, TM, U["squad"],
     "Senza asola: allunga a mano il foro verso il piano con una punta Ø6"),
    ("Ferramenta telaio", "Viti legno assortite (5×60, 5×70, 4×35, 4×50, 4,5×25)", "rinforzo, listelli, ripiani, squadrette", 1, "set", 15.00, S, False, TM, "", ""),
    ("Ferramenta telaio", "Staffa angolare con nervatura 90×90×65 zincata", "fissaggio a muro (G1P e G4P)", 4, "pz", 2.89, V, False, OB, U["staffa"], ""),
    ("Ferramenta telaio", "Tasselli Ø8 per il tipo di muro", "fissaggio a muro", 1, "conf", 3.00, S, False, TM, "", ""),
    ("Zona ricarica", "Piastre acciaio tagliate e forate (base 180×180×12, contropiastra 220×220×6, 3 piastre attrezzo)",
     "attacco presse e morsa", 1, "lotto", 90.00, S, False, "officina/carpenteria", "", "Chiedi un preventivo a un'officina o a un servizio di taglio laser"),
    ("Zona ricarica", "Vite TSPEI M10×110 + dado autobloccante", "base, 2 posteriori", 2, "pz", 1.50, S, False, FO, "", "Non trovate a scaffale"),
    ("Zona ricarica", "Vite TSPEI M10×200 + dado autobloccante", "base, 2 anteriori", 2, "pz", 2.00, S, False, FO, "", "Non trovate a scaffale"),
    ("Zona ricarica", "Dadi autobloccanti M10 inox, conf. 10 pz", "sotto la base", 1, "conf", 2.40, V, False, TM, U["dadiaut"], ""),
    ("Zona ricarica", "Rondelle piane larghe M10 (Ø29,5), conf. 8 pz", "sotto la traversa e la contropiastra", 1, "conf", 2.50, V, False, LM, U["rondlarghe"], ""),
    ("Zona ricarica", "Manopole a stella M12 + 2 spine Ø10 h8", "bloccaggio piastre", 1, "kit", 10.00, S, False, FO, "", ""),
    ("Elettrico", "Canalina di distribuzione Arnocanali 60×40, 2 scomparti", "una per campata (da tagliare)", 3, "pz", 8.55, V, False, TM, U["canalina"],
     "Lunghezza non indicata: verifica. Alternativa specifica sottoscrivania: LM 76 cm 36,95 € (online)"),
    ("Elettrico", "Bocchette passacavi Ø60 PVC nero, conf. 2 pz", "piano", 1, "conf", 1.05, V, False, TM, U["passacavi"], ""),
    ("Elettrico", "Ciabatta 6 prese con interruttore", "campata seduta", 1, "pz", 20.00, S, False, TM, "", ""),
    ("Finitura", "Olio Adler facile applicazione interno 750 ml", "piano di lavoro", 1, "pz", 27.90, V, False, TM, U["olio"], ""),
    ("Finitura", "Impregnante legno all'acqua Cipir 0,75 l incolore", "telaio e ripiani OSB", 1, "pz", 6.70, V, False, TM, U["impr"], ""),
    ("Finitura", "Abrasivi 80–120–180", "levigatura", 1, "set", 8.00, S, False, TM, "", ""),
]
distinta(ws, "Modulo 1 – Tavolo base: distinta e costi",
         f"Prezzi verificati sui siti di Tecnomat, OBI e Leroy Merlin il {DATA} (possono variare per negozio). Celle gialle = stime.",
         base_items, ["Legno", "Ferramenta telaio", "Zona ricarica", "Elettrico", "Finitura"],
         ["Tecnomat: taglio legno in negozio, 1° taglio gratis poi 1 € a taglio (solo per pannelli e listelli comprati lì).",
          "I prezzi Tecnomat sono del negozio predefinito del sito (Roma Prenestina): controlla quello dove andrai.",
          "I 4 dadi a tamburo della predisposizione alzatina vanno inseriti durante il montaggio del tavolo."])

wt = wb.create_sheet("Confronto travi")
wt["A1"] = "Confronto travi per il telaio"; wt["A1"].font = H1
header(wt, 4, ["Trave", "Sezione", "Lungh. barra (m)", "Prezzo barra", "Prezzo al metro", "Metri necessari", "Rigidezza gamba (8×8=1)",
               "", "Barre necessarie", "Costo telaio"], [30, 11, 11, 12, 13, 12, 14, 2, 11, 13])
for i, r in enumerate(cf["results"]):
    row = 5 + i
    setc(wt, f"A{row}", r["name"], BOLD if r["kvh"] and r["b"] == 80 else NORM)
    setc(wt, f"B{row}", f'{r["b"]}×{r["h"]}'); setc(wt, f"C{row}", r["stock"] / 1000, BLUE, "0.0")
    setc(wt, f"D{row}", r["price"], BLUE, EUR); setc(wt, f"E{row}", f"=D{row}/C{row}", NORM, EURM)
    setc(wt, f"F{row}", r["need_mm"] / 1000, BLUE, "0.00"); setc(wt, f"G{row}", round(r["I_rel"], 1), BLUE, "0.0")
    setc(wt, f"I{row}", len(r["bars"]), BLUE, "0"); setc(wt, f"J{row}", f"=I{row}*D{row}", NORM, EUR)
    if r["kvh"] and r["b"] == 80:
        for j in range(1, 11): wt.cell(row=row, column=j).fill = SEL
assert cf["results"][3]["name"].startswith("KVH giuntato 80")
wt["A11"] = "Riga verde = trave scelta (Tecnomat). Lamellare GL24h 80×120×4000 Tecnomat 39,90 €; Leroy Merlin 74,90–79,90 €; OBI non ha travi 80×120."
wt["A11"].font = NOTE

wr = wb.create_sheet("Confronto ripiani")
wr["A1"] = "Ripiani bassi e rinforzo ricarica: materiali a confronto"; wr["A1"].font = H1
header(wr, 4, ["Materiale", "Negozio", "Formato", "Pezzi", "Prezzo", "Costo", "Note"], [32, 14, 18, 8, 12, 12, 70])
opts = [("OSB/3 18 (SCELTO)", TM, "2500×1250", 1, 27.50, "Un pannello per 3 ripiani + rinforzo in 2 strati; avanzano ~1,8 m²"),
        ("OSB/3 18", LM, "2500×1250", 1, 30.50, ""), ("OSB/3 18", OB, "2500×1250", 1, 31.49, ""),
        ("Lamellare abete 40", TM, "2000×800", 1, 79.00, "Ripiani più spessi e belli, 2,9 volte l'OSB"),
        ("Tavola lamellare 14", OB, "800×200", 14, 8.19, "8 per i ripiani + 6 per il rinforzo: la più cara")]
for i, (m, shop, fmt, q, pr, note) in enumerate(opts):
    row = 5 + i
    setc(wr, f"A{row}", m, BOLD if i == 0 else NORM); setc(wr, f"B{row}", shop); setc(wr, f"C{row}", fmt)
    setc(wr, f"D{row}", q, BLUE, "0"); setc(wr, f"E{row}", pr, BLUE, EUR); setc(wr, f"F{row}", f"=D{row}*E{row}", NORM, EUR)
    setc(wr, f"G{row}", note, wrap=True)
    if i == 0:
        for j in range(1, 8): wr.cell(row=row, column=j).fill = SEL

wc = wb.create_sheet("Taglio travi")
wc["A1"] = "Taglio travi KVH 80×120 (barre da 4 m)"; wc["A1"].font = H1
header(wc, 3, ["Barra", "Tagli (mm)", "Usato (mm)", "Avanzo (mm)"], [10, 52, 12, 12])
for i, bar in enumerate(cf["kvh_bars"], 1):
    row = 3 + i
    setc(wc, f"A{row}", f"Barra {i}"); setc(wc, f"B{row}", " + ".join(map(str, bar)))
    setc(wc, f"C{row}", sum(bar) + 4 * len(bar) + 10, BLUE, "#,##0"); setc(wc, f"D{row}", f"=4000-C{row}", NORM, "#,##0")
wc[f"A{5 + len(cf['kvh_bars'])}"] = "I 4 pezzi da 250 sono i montanti dell'alzatina: tagliali e conservali. Tecnomat taglia in negozio (1 € a taglio dopo il primo)."
wc[f"A{5 + len(cf['kvh_bars'])}"].font = NOTE
wb.save("tavolo_base_costi.xlsx")

# ============================================================== 2. CASSETTIERA
cas = [p for p in parts if p["module"] == "cassettiera"]
pieces = {"multistrato 18": [], "multistrato 15": [], "HDF 4": []}
for p in cas:
    a, b, c = dims(p)
    if "(scatola)" in p["name"]:
        x0, y0, z0, x1, y1, z1 = p["box"]; wd, l, h = round(x1 - x0, 1), round(y1 - y0), round(z1 - z0)
        tag = "appeso" if "appeso" in p["name"] else p["name"].split(" (")[0].lower()
        pieces["multistrato 15"] += [(f"{tag}: fianco", l, h)] * 2 + [(f"{tag}: testata", round(wd - 30), h)] * 2
        pieces["HDF 4"] += [(f"{tag}: fondo", l, round(wd))]
    else:
        pieces[p["mat"]].append((p["name"], a, b))
json.dump(pieces, open("cassettiera_pezzi.json", "w"), indent=1)
wb = Workbook(); ws = wb.active; ws.title = "Distinta costi"
cas_items = [
    ("Legno", "Multistrato pioppo 18 mm 1200×600", "scatola e frontali (2 pannelli, schema nel manuale)", 2, "pz", 17.80, V, False, TM, U["pio18"],
     "Il pannello intero da 18 non c'è a Tecnomat; LM al m² 59 €/m²"),
    ("Legno", "Multistrato pioppo 15 mm 1200×600", "scatole dei 4 cassetti", 2, "pz", 15.90, V, False, TM, U["pio15"], ""),
    ("Legno", "Fibra grezza (faesite) 3 mm 2500×1250", "fondi dei cassetti e schienale", 1, "pz", 16.00, V, False, TM, U["fibra3"],
     "3 mm invece di 4: va bene per carichi fino a ~15 kg per cassetto"),
    ("Ferramenta", "Coppia guide cassetto estrazione totale H45 L 500", "3 cassettiera + 1 appeso", 4, "cp", 10.90, V, False, TM, U["guide"],
     "Portata non indicata; LM Cisne-45 a sfere 50 kg a 10,90 €"),
    ("Ferramenta", "Maniglie", "", 4, "pz", 4.00, S, False, TM, "", ""),
    ("Ferramenta", "Viti 4×45, 3,5×30, 3,5×25, 3×16, 4×16, 4×30, 4×50", "scatola, cassetti, frontali, guide, fissaggio", 1, "set", 10.00, S, False, TM, "", ""),
    ("Ferramenta", "Collante vinilico D3 Tover 500 ml", "giunti", 1, "pz", 4.95, V, False, TM, U["colla"], ""),
    ("Finitura", "Impregnante legno all'acqua Cipir 0,75 l", "scatola e frontali", 1, "pz", 6.70, V, False, TM, U["impr"], ""),
]
distinta(ws, "Modulo 2 – Cassettiera: distinta e costi",
         f"Cassettiera fissa per la campata sinistra (3 cassetti) + cassetto appeso. Prezzi verificati il {DATA}; celle gialle = stime.",
         cas_items, ["Legno", "Ferramenta", "Finitura"],
         ["Alternativa economica: scatola in OSB 18 dall'avanzo del pannello del tavolo base (0 €).",
          "Tecnomat taglia i pannelli in negozio: porta il foglio 'Taglio pannelli'."])
wp = wb.create_sheet("Taglio pannelli")
wp["A1"] = "Pezzi da tagliare"; wp["A1"].font = H1
header(wp, 3, ["Materiale", "Pezzo", "Lunghezza", "Larghezza", "Q.tà"], [18, 34, 12, 12, 8])
row = 4
for mat, lst in pieces.items():
    for (n, a, b), q in Counter(lst).items():
        setc(wp, f"A{row}", mat); setc(wp, f"B{row}", n); setc(wp, f"C{row}", a, BLUE, "0"); setc(wp, f"D{row}", b, BLUE, "0")
        setc(wp, f"E{row}", q, BLUE, "0"); row += 1
wb.save("cassettiera_costi.xlsx")

# ============================================================== 3. ALZATINA
wb = Workbook(); ws = wb.active; ws.title = "Distinta costi"
alz_items = [
    ("Legno", "Montanti KVH 80×120×250", "tagliati con le travi del tavolo base", 4, "pz", 0, B, False, TM, "", ""),
    ("Legno", "Tavola in abete lamellare 1000×200×18", "mensola: tagliate a 620, 780 e 600", 3, "pz", 4.50, V, False, TM, U["tav18"],
     "OBI 800×200×18 8,79 €; OBI 800×200×14 8,19 €"),
    ("Legno", "Listello abete grezzo 30×40×2000", "rinforzo sotto il bordo anteriore", 1, "pz", 3.10, V, False, TM, U["list2"],
     "Visibile dal basso: levigalo. Piallato 30×40×2000 da OBI a 3,99 €"),
    ("Ferramenta", "Barra filettata M10×1000 zincata", "4 spezzoni da 390", 2, "pz", 1.65, V, False, TM, U["barra"], ""),
    ("Ferramenta", "Dado cieco M10 + rondella Ø30", "sopra la mensola", 4, "set", 0.80, S, False, TM, "", ""),
    ("Ferramenta", "Dadi a tamburo M10 nelle gambe posteriori", "già montati con il tavolo base", 4, "pz", 0, B, False, FO, "", ""),
    ("Ferramenta", "Viti 4×50", "listello sotto le tavole", 1, "conf", 3.00, S, False, TM, "", ""),
    ("Ferramenta", "Collante vinilico D3 Tover 500 ml", "listello", 1, "pz", 4.95, V, False, TM, U["colla"], "Se avanza dal tavolo base: 0 €"),
    ("Ferramenta", "Staffa angolare con nervatura 90×90×65 zincata", "fissaggio a muro su M1 e M4", 2, "pz", 2.89, V, False, OB, U["staffa"], ""),
    ("Ferramenta", "Tasselli Ø8", "fissaggio a muro", 1, "conf", 2.00, S, False, TM, "", ""),
    ("Finitura", "Impregnante (quota)", "mensola e montanti", 1, "pz", 3.00, S, False, TM, "", ""),
]
distinta(ws, "Modulo 3 – Alzatina con mensola: distinta e costi",
         f"Si monta sul tavolo base usando le predisposizioni. Prezzi verificati il {DATA}; celle gialle = stime.",
         alz_items, ["Legno", "Ferramenta", "Finitura"],
         ["Se il tavolo base è già montato senza i montanti, serve 1 m di KVH 80×120: una barra da 4 m costa 29,90 € (Tecnomat)."])
wb.save("alzatina_costi.xlsx")

# ============================================================== CONFRONTO NEGOZI
wb = Workbook(); ws = wb.active; ws.title = "Prezzi per negozio"
ws["A1"] = f"Confronto prezzi Tecnomat · OBI · Leroy Merlin ({DATA})"; ws["A1"].font = H1
ws["A2"] = "Solo prezzi letti sulle pagine dei siti. LM = Leroy Merlin; MP = venditore terzo (marketplace) con spedizione a parte."; ws["A2"].font = NOTE
header(ws, 4, ["Voce", "Negozio", "Prodotto", "Misure", "Prezzo", "Conf.", "Link", "Note"], [24, 12, 50, 18, 11, 9, 7, 46])
R = [
 ("Trave 80×120", TM, "Trave KVH giuntato 80×120×4000", "80×120×4000", 29.90, "1", U["kvh"], "scelta"),
 ("Trave 80×120", TM, "Trave abete lamellare GL24h 80×120×4000", "80×120×4000", 39.90, "1", "https://www.tecnomat.it/it/trave-lamellare-abete-8x12x400-cm-10074468/", ""),
 ("Trave 80×120", LM, "Trave lamellare abete industriale 80×120×4000", "80×120×4000", 74.90, "1", "https://www.leroymerlin.it/prodotti/trave-lamellare-abete-industriale-l-400-x-p-12-cm-sp-80-mm-35613963.html", ""),
 ("Trave 80×120", LM, "Trave lamellare abete piallato 80×120×4000", "80×120×4000", 79.90, "1", "https://www.leroymerlin.it/prodotti/trave-lamellare-in-abete-piallato-l-400-x-p-8-x-sp-12-cm-35282156.html", ""),
 ("Trave (alternativa)", OB, "Travetto abete massello angoli arrotondati 90×90×3000", "90×90×3000", 41.99, "1", "https://www.obi-italia.it/p/2872109", "OBI non ha 80×120"),
 ("Piano 2000×800×40", TM, "Tavola abete lamellare 200×80 sp. 40", "2000×800×40", 79.00, "1", U["lam40"], "scelta"),
 ("Piano 2000×800×40", LM, "Tavola lamellare abete 40 mm 80×200 (MP Cardigliano)", "2000×800×40", 147.45, "1", "https://www.leroymerlin.it/prodotti/tavola-pannello-lamellare-abete-spessore-mm-40-80-cm-200-cm-84286356.html", "solo online"),
 ("Piano 2000×800 (28 mm)", TM, "Tavola abete lamellare 200×80 sp. 28", "2000×800×28", 49.90, "1", "https://www.tecnomat.it/it/tavola-abete-lam-200x80-cm-spess-28-10074974/", ""),
 ("Piano 2000×800 (27 mm)", LM, "Tavola lamellare abete 80×200 sp. 27", "2000×800×27", 69.90, "1", "https://www.leroymerlin.it/prodotti/tavola-lamellare-in-legno-di-abete-l-80-x-200-cm-x-sp-27-mm-86909783.html", ""),
 ("OSB/3 18", TM, "OSB 3 di pioppo 18 mm 250×125", "2500×1250×18", 27.50, "1", U["osb"], "scelta"),
 ("OSB/3 18", LM, "Pannello OSB 3 250×125 sp. 18", "2500×1250×18", 30.50, "1", "https://www.leroymerlin.it/prodotti/pannello-osb-3-l-250-x-h-125-cm-sp-18-mm-naturale-89126386.html", ""),
 ("OSB/3 18", OB, "Pannello OSB/3 2500×1250×18", "2500×1250×18", 31.49, "1", "https://www.obi-italia.it/p/2099109", ""),
 ("Multistrato pioppo 18", TM, "Multistrato pioppo 18 mm 120×60", "1200×600×18", 17.80, "1", U["pio18"], "scelta"),
 ("Multistrato pioppo 18", LM, "Compensato pioppo 18 mm al taglio", "al m²", 59.00, "m²", "https://www.leroymerlin.it/prodotti/pannello-compensato-pioppo-l-250-x-h-122-cm-sp-18-mm-naturale-su-misura-pioppo-33000702.html", ""),
 ("Multistrato pioppo 15", TM, "Multistrato pioppo 15 mm 120×60", "1200×600×15", 15.90, "1", U["pio15"], "scelta"),
 ("Multistrato pioppo 15", TM, "Multistrato pioppo 15 mm 244×122", "2440×1220×15", 47.60, "1", "https://www.tecnomat.it/it/pioppo-spess-15-mm-244x122-cm-25022700/", ""),
 ("Multistrato pioppo 15", OB, "Pannello pioppo multistrato 15×1200×900", "1200×900×15", 49.99, "1", "https://www.obi-italia.it/p/5533476", ""),
 ("Multistrato pioppo 15", LM, "Multistrato pioppo 250×122 sp. 15", "2500×1220×15", 80.50, "1", "https://www.leroymerlin.it/prodotti/pannello-multistrato-in-pioppo-l-250-x-h-122-cm-sp-15-mm-naturale-89185626.html", ""),
 ("Fondi cassetti", TM, "Fibra grezza 3 mm 250×125", "2500×1250×3", 16.00, "1", U["fibra3"], "scelta"),
 ("Fondi cassetti", OB, "MDF pretagliato 1010×700×4", "1010×700×4", 9.19, "1", "https://www.obi-italia.it/p/3809142", ""),
 ("Tavola 200 mm larga", TM, "Tavola abete lamellare 100×20 sp. 18", "1000×200×18", 4.50, "1", U["tav18"], "scelta (mensola)"),
 ("Tavola 200 mm larga", OB, "Tavola lamellare abete 800×200×18", "800×200×18", 8.79, "1", "https://www.obi-italia.it/p/2190791", ""),
 ("Tavola 200 mm larga", OB, "Tavola lamellare abete 800×200×14", "800×200×14", 8.19, "1", "https://www.obi-italia.it/p/2423135", ""),
 ("Listello 30×40", TM, "Listello abete grezzo 2 m 30×40", "30×40×2000", 3.10, "1", U["list2"], "scelta"),
 ("Listello 30×40", OB, "Listello abete piallato 4 lati 30×40×2000", "30×40×2000", 3.99, "1", "https://www.obi-italia.it/p/6191670", ""),
 ("Vite M10 lunga", OB, "Vite testa cilindrica zincata M10×180", "M10×180", 1.39, "1", U["m10x180"], "scelta (B1)"),
 ("Vite M10 lunga", OB, "Viti testa cilindrica M10×160", "M10×160", 13.09, "10", U["m10x160"], "scelta (B2, da accorciare)"),
 ("Vite M10 lunga", OB, "Vite testa esagonale M10×140", "M10×140", 3.99, "5", "https://www.obi-italia.it/p/6271043", "testa esagonale: lamatura Ø30"),
 ("Vite M10 lunga", TM, "Bulloni testa tonda collo quadro M10×160 Würth", "M10×160", 19.45, "25", "https://www.tecnomat.it/it/bulloni-t-tonda-quadra-zinc-m10x160-25pz-10049377/", ""),
 ("Vite M10 lunga", LM, "Vite testa esagonale DIN 931 M10×140 (MP)", "M10×140", 35.31, "50", "https://www.leroymerlin.it/prodotti/vite-con-filetto-metrico-a-testa-esagonale-8-8-filetto-medio-din-931-zincato-m10-x-140-scatola-da-50-unita-86447822.html", "spedizione 12,50 €"),
 ("Barra filettata M10 1 m", TM, "Barra filettata M10×1 m zincata", "M10×1000", 1.65, "1", U["barra"], "scelta"),
 ("Barra filettata M10 1 m", LM, "Barra filettata SUKI M10×1000", "M10×1000", 1.75, "1", "https://www.leroymerlin.it/prodotti/barra-filettata-suki-in-acciaio-m10-l-1000-mm-35737030.html", ""),
 ("Barra filettata M10 1 m", OB, "Barra filettata zincata M10×1000", "M10×1000", 3.19, "1", "https://www.obi-italia.it/p/4328456", ""),
 ("Spine Ø12", OB, "Spine in faggio zigrinato Ø12", "Ø12", 3.69, "25", U["spine"], "scelta"),
 ("Spine Ø12", TM, "Tondo faggio zigrinato Ø12×1000", "Ø12×1000", 0.80, "1", "https://www.tecnomat.it/it/tondo-faggio-zigrinato-o-12-mm-h-100-cm-10018430/", "da tagliare"),
 ("Guide cassetto 500", TM, "Coppia guide estrazione totale H45 L500", "500", 10.90, "coppia", U["guide"], "scelta"),
 ("Guide cassetto 500", LM, "Guide Cisne-45 a sfere 50 kg L500 (MP)", "500", 10.90, "?", "https://www.leroymerlin.it/prodotti/guide-per-cassetto-cisne-45-a-sfere-h45-mm-estrazione-totale-portata-50-kg-lunghezza-500-mm-chiusura-standard-evotrade-84639730.html", ""),
 ("Guide cassetto 500", OB, "Guida a sfere Hettich Silent System 50 cm", "500", 37.99, "?", "https://www.obi-italia.it/p/4503207", ""),
 ("Staffa angolare 90×90", OB, "Staffa angolare con nervatura 90×90×65", "90×90×65", 2.89, "1", U["staffa"], "scelta"),
 ("Staffa angolare 90×90", TM, "Squadretta di giunzione angolare 90×90×65", "90×90×65", 0.80, "1", "https://www.tecnomat.it/it/squadretta-di-giunz-angol-90x90x65mm-10062291/", "nervatura non indicata"),
 ("Passacavi Ø60", TM, "Bocchette passacavi Ø60 PVC nero", "Ø60", 1.05, "2", U["passacavi"], "scelta"),
 ("Passacavi Ø60", LM, "Passacavo Hettich Ø60 grigio", "Ø60", 0.99, "1", "https://www.leroymerlin.it/prodotti/passacavo-hettich-grigio-argento-60-mm-43826482.html", ""),
 ("Canalina", TM, "Canalina Arnocanali 60×40 2 scomparti", "60×40", 8.55, "1", U["canalina"], "scelta (lunghezza da verificare)"),
 ("Canalina", LM, "Canalina per scrivanie e tavoli 76 cm (MP)", "760×185×100", 36.95, "1", "https://www.leroymerlin.it/prodotti/canalina-per-scrivanie-e-tavoli-76-cm-84891827.html", "spedizione da 15,95 €"),
]
for i, (v, shop, prod, mis, pr, conf, url, note) in enumerate(R):
    row = 5 + i
    setc(ws, f"A{row}", v); setc(ws, f"B{row}", shop); setc(ws, f"C{row}", prod, wrap=True); setc(ws, f"D{row}", mis)
    setc(ws, f"E{row}", pr, BLUE, EUR); setc(ws, f"F{row}", conf); link(ws, f"G{row}", url); setc(ws, f"H{row}", note, wrap=True)
    if note.startswith("scelta"):
        for j in range(1, 9): ws.cell(row=row, column=j).fill = SEL
ws.freeze_panes = "A5"
ws[f"A{6 + len(R)}"] = "Non trovati in nessuno dei tre negozi: dadi a tamburo M10, viti TCEI M10×140/180 dichiarate, viti TSPEI M10 lunghe, inserti M10 a vite (solo MP)."
ws[f"A{6 + len(R)}"].font = NOTE

wj = wb.create_sheet("Giunzioni alternative")
wj["A1"] = "Sistemi di giunzione traverse–gambe: costo per 34 nodi"; wj["A1"].font = H1
header(wj, 3, ["Sistema", "Componente", "Negozio", "Q.tà", "Prezzo", "Totale", "Note"], [30, 46, 14, 7, 11, 12, 60])
J = [
 ("A · Bulloni + dadi a tamburo (attuale)", "Vite M10×180 testa cilindrica", OB, 16, 1.39, "smontabile, rigido; fori precisi con dima"),
 ("A · Bulloni + dadi a tamburo (attuale)", "Viti M10×160 conf. 10 (accorciate)", OB, 1, 13.09, ""),
 ("A · Bulloni + dadi a tamburo (attuale)", "Dadi a tamburo M10 (stima, online)", FO, 38, 0.80, "unico pezzo non trovato a scaffale"),
 ("A · Bulloni + dadi a tamburo (attuale)", "Barra filettata M10 1 m", TM, 1, 1.65, ""),
 ("A · Bulloni + dadi a tamburo (attuale)", "Spine Ø12 conf. 25", OB, 3, 3.69, ""),
 ("B · Viti strutturali", "Viti Rothoblaas TPS 8×200 conf. 100", TM, 1, 41.90, "2 viti per nodo (68), inclinate; si smonta poche volte"),
 ("B · Viti strutturali", "Spine Ø12 conf. 25", OB, 3, 3.69, "antirotazione"),
 ("C · Staffe metalliche", "Staffa angolare con nervatura 90×90×65 (conf. 20 a 29,99 €)", OB, 4, 29.99, "2 per nodo (68 → 4 conf.); si vedono, meno rigido"),
 ("C · Staffe metalliche", "Viti per staffe (stima)", TM, 1, 12.00, ""),
]
for i, (sysn, comp, shop, q, pr, note) in enumerate(J):
    row = 4 + i
    setc(wj, f"A{row}", sysn); setc(wj, f"B{row}", comp, wrap=True); setc(wj, f"C{row}", shop)
    setc(wj, f"D{row}", q, BLUE, "0"); setc(wj, f"E{row}", pr, BLUE, EUR); setc(wj, f"F{row}", f"=D{row}*E{row}", NORM, EUR)
    setc(wj, f"G{row}", note, wrap=True)
r0 = 5 + len(J)
for k, sysn in enumerate(["A · Bulloni + dadi a tamburo (attuale)", "B · Viti strutturali", "C · Staffe metalliche"]):
    setc(wj, f"E{r0 + k}", f"Totale {sysn[0]}", BOLD, border=False)
    setc(wj, f"F{r0 + k}", f'=SUMIFS(F4:F{3 + len(J)},A4:A{3 + len(J)},"{sysn}")', BOLD, EUR)
wb.save("confronto_prezzi_negozi.xlsx")
print("ok")
